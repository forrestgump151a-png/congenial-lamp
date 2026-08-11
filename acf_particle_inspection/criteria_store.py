#!/usr/bin/env python3
"""
ACF(異方性導電膜)粒子検査 規格値ストア

品名ごとの粒子検査規格(一次粒子径・平均粒子径・粒子面積率の許容範囲)を
CSV/Excel から読み込み、品名検索できるようにする。
"""

from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

DATA_DIR = Path(__file__).parent / "data"
CRITERIA_PATH = DATA_DIR / "criteria.csv"

# CSV/Excel に必須の列名
REQUIRED_COLUMNS = [
    "品名",
    "一次粒子径_下限", "一次粒子径_上限", "一次粒子径_単位",
    "平均粒子径_下限", "平均粒子径_上限", "平均粒子径_単位",
    "粒子面積率_下限", "粒子面積率_上限", "粒子面積率_単位",
]
OPTIONAL_COLUMNS = ["備考", "更新日"]
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


@dataclass
class Range:
    lower: float | None
    upper: float | None
    unit: str = ""

    def as_text(self) -> str:
        if self.lower is None and self.upper is None:
            return "-"
        if self.lower is not None and self.upper is not None:
            return f"{_fmt(self.lower)} ~ {_fmt(self.upper)} {self.unit}".strip()
        if self.lower is not None:
            return f"{_fmt(self.lower)} {self.unit} 以上".strip()
        return f"{_fmt(self.upper)} {self.unit} 以下".strip()

    def contains(self, value: float) -> bool | None:
        if self.lower is None and self.upper is None:
            return None
        if self.lower is not None and value < self.lower:
            return False
        if self.upper is not None and value > self.upper:
            return False
        return True


def _fmt(v: float) -> str:
    s = f"{v:.4f}".rstrip("0").rstrip(".")
    return s if s else "0"


@dataclass
class ProductCriteria:
    name: str
    primary_diameter: Range
    average_diameter: Range
    area_ratio: Range
    note: str = ""
    updated_at: str = ""
    search_key: str = field(default="", repr=False)

    def __post_init__(self):
        self.search_key = normalize(self.name)


def normalize(text: str) -> str:
    """全角/半角・大小文字・前後空白を吸収した検索用キーを作る"""
    if text is None:
        return ""
    text = unicodedata.normalize("NFKC", str(text))
    return text.strip().lower()


def _to_float(raw: str) -> float | None:
    raw = (raw or "").strip()
    if raw == "":
        return None
    raw = unicodedata.normalize("NFKC", raw)
    try:
        return float(raw)
    except ValueError:
        return None


class CriteriaFileError(ValueError):
    pass


def _normalize_header(h: str) -> str:
    return unicodedata.normalize("NFKC", (h or "")).strip()


def parse_rows(rows: list[dict]) -> list[ProductCriteria]:
    items: list[ProductCriteria] = []
    for i, row in enumerate(rows, start=2):  # 2: ヘッダーが1行目
        row = {_normalize_header(k): v for k, v in row.items()}
        name = (row.get("品名") or "").strip()
        if not name:
            continue
        try:
            item = ProductCriteria(
                name=name,
                primary_diameter=Range(
                    _to_float(row.get("一次粒子径_下限", "")),
                    _to_float(row.get("一次粒子径_上限", "")),
                    (row.get("一次粒子径_単位") or "").strip(),
                ),
                average_diameter=Range(
                    _to_float(row.get("平均粒子径_下限", "")),
                    _to_float(row.get("平均粒子径_上限", "")),
                    (row.get("平均粒子径_単位") or "").strip(),
                ),
                area_ratio=Range(
                    _to_float(row.get("粒子面積率_下限", "")),
                    _to_float(row.get("粒子面積率_上限", "")),
                    (row.get("粒子面積率_単位") or "").strip(),
                ),
                note=(row.get("備考") or "").strip(),
                updated_at=(row.get("更新日") or "").strip(),
            )
        except Exception as e:
            raise CriteriaFileError(f"{i}行目の読み込みに失敗しました: {e}") from e
        items.append(item)
    return items


def load_csv_text(text: str) -> list[ProductCriteria]:
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise CriteriaFileError("CSVにヘッダー行がありません")
    headers = {_normalize_header(h) for h in reader.fieldnames}
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise CriteriaFileError(
            "必須列が不足しています: " + ", ".join(missing)
        )
    return parse_rows(list(reader))


def load_xlsx_bytes(data: bytes) -> list[ProductCriteria]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise CriteriaFileError(
            "Excelファイルを読むには openpyxl が必要です (pip install openpyxl)"
        ) from e
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [_normalize_header(str(h) if h is not None else "") for h in next(rows_iter)]
    except StopIteration:
        raise CriteriaFileError("Excelにデータがありません")
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise CriteriaFileError("必須列が不足しています: " + ", ".join(missing))
    dict_rows = []
    for row in rows_iter:
        d = {header[i]: ("" if v is None else str(v)) for i, v in enumerate(row) if i < len(header)}
        dict_rows.append(d)
    return parse_rows(dict_rows)


def load_criteria(path: Path = CRITERIA_PATH) -> list[ProductCriteria]:
    if not path.exists():
        return []
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return load_xlsx_bytes(path.read_bytes())
    return load_csv_text(path.read_text(encoding="utf-8-sig"))


def search(items: list[ProductCriteria], query: str, exact: bool = False) -> list[ProductCriteria]:
    key = normalize(query)
    if not key:
        return list(items)
    if exact:
        return [it for it in items if it.search_key == key]
    return [it for it in items if key in it.search_key]
